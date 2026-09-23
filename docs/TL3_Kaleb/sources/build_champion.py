#!/usr/bin/env python3
"""Build the Feature 4 Champion document and matching diagrams."""

from __future__ import annotations

import os
from pathlib import Path

os.environ.setdefault("MPLCONFIGDIR", "/tmp/mplconfig")

import matplotlib.pyplot as plt
from matplotlib.patches import (
    Circle,
    Ellipse,
    FancyArrowPatch,
    FancyBboxPatch,
    Rectangle,
)
from matplotlib.lines import Line2D
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT_DOCX = ROOT / "Champion_Upgrade_and_Loot_System.docx"
IMG = ROOT / "general_diagrams"
IMG.mkdir(exist_ok=True)

NAVY = "#2F3E4E"
NAVY_RGB = RGBColor(0x1E, 0x27, 0x61)
PEACH = "#F3D2B5"
PEACH_EDGE = "#C4A882"
GRAY = "#E8E8E8"
WHITE = "#FFFFFF"
GOLD = "#C1962A"
GOLD_LIGHT = "#F3E4C4"
TITLE_C = "#1E2761"
TEXT = "#2C3E50"
MUTED = "#5A6A7A"


def save(fig, name: str) -> Path:
    path = IMG / name
    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white", pad_inches=0.18)
    plt.close(fig)
    return path


def new_fig(w=11.2, h=6.6, equal=True, xmax=100, ymax=100):
    fig, ax = plt.subplots(figsize=(w, h))
    ax.set_xlim(0, xmax)
    ax.set_ylim(0, ymax)
    if equal:
        ax.set_aspect("equal")
    ax.axis("off")
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    return fig, ax


def text(ax, x, y, s, size=10, weight="normal", color=TEXT, ha="center", va="center", style="normal"):
    ax.text(
        x,
        y,
        s,
        fontsize=size,
        fontweight=weight,
        color=color,
        ha=ha,
        va=va,
        fontstyle=style,
        zorder=6,
        clip_on=False,
    )


def rounded(ax, x, y, w, h, facecolor=WHITE, edge=NAVY, lw=1.6, radius=0.18):
    p = FancyBboxPatch(
        (x, y),
        w,
        h,
        boxstyle=f"round,pad=0.15,rounding_size={min(w, h) * radius}",
        facecolor=facecolor,
        edgecolor=edge,
        linewidth=lw,
        zorder=3,
        mutation_aspect=1,
    )
    ax.add_patch(p)
    return p


def rect(ax, x, y, w, h, facecolor=WHITE, edge=NAVY, lw=1.6, z=3):
    p = Rectangle((x, y), w, h, facecolor=facecolor, edgecolor=edge, linewidth=lw, zorder=z)
    ax.add_patch(p)
    return p


def gane_process(ax, cx, cy, w, h, pid, name, fill=WHITE):
    x, y = cx - w / 2, cy - h / 2
    rounded(ax, x, y, w, h, facecolor=fill, radius=0.12)
    header = h * 0.38
    ax.plot([x + 0.6, x + w - 0.6], [y + h - header, y + h - header], color=NAVY, lw=1.15, zorder=4)
    text(ax, cx, y + h - header / 2, pid, size=9.5, weight="bold")
    text(ax, cx, y + (h - header) / 2 - 0.1, name, size=9)


def data_store(ax, cx, cy, w, h, did, name):
    """Gane–Sarson store: open on the right, left-hand identifier band."""
    x, y = cx - w / 2, cy - h / 2
    ax.add_patch(Rectangle((x, y), w, h, facecolor=WHITE, edgecolor="none", linewidth=0, zorder=3))
    split = w * 0.22
    ax.plot([x, x + w], [y + h, y + h], color=NAVY, lw=1.55, zorder=4, solid_capstyle="butt")
    ax.plot([x, x + w], [y, y], color=NAVY, lw=1.55, zorder=4, solid_capstyle="butt")
    ax.plot([x, x], [y, y + h], color=NAVY, lw=1.55, zorder=4, solid_capstyle="butt")
    ax.plot([x + split, x + split], [y, y + h], color=NAVY, lw=1.55, zorder=4, solid_capstyle="butt")
    text(ax, x + split / 2, cy, did, size=9, weight="bold")
    text(ax, x + split + (w - split) / 2, cy, name, size=9)


def entity(ax, cx, cy, w, h, name):
    x, y = cx - w / 2, cy - h / 2
    rect(ax, x, y, w, h, facecolor=GRAY, lw=1.7)
    text(ax, cx, cy, name, size=10, weight="bold")


def arrow(ax, p1, p2, color=NAVY, lw=1.25, rad=0.0, style="-|>"):
    ax.add_patch(
        FancyArrowPatch(
            p1,
            p2,
            arrowstyle=style,
            mutation_scale=11,
            linewidth=lw,
            color=color,
            connectionstyle=f"arc3,rad={rad}",
            zorder=2,
            shrinkA=0,
            shrinkB=0,
        )
    )


def assoc(ax, p1, p2, color=NAVY, lw=1.5):
    """Communicates: solid line, no arrowheads."""
    ax.plot([p1[0], p2[0]], [p1[1], p2[1]], color=color, lw=lw, zorder=2, solid_capstyle="round")


def open_head(ax, p1, p2, size=2.55, color=NAVY, lw=1.35):
    """Open chevron at p2 (include / extend). Two strokes, not a filled triangle."""
    x1, y1 = p1
    x2, y2 = p2
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1e-6)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    wing = size
    spread = size * 0.5
    left = (x2 - ux * wing + px * spread, y2 - uy * wing + py * spread)
    right = (x2 - ux * wing - px * spread, y2 - uy * wing - py * spread)
    for wx, wy in (left, right):
        ax.plot(
            [wx, x2],
            [wy, y2],
            color=color,
            lw=lw,
            solid_capstyle="butt",
            zorder=5,
        )


def dashed(ax, p1, p2, color=NAVY, lw=1.35, head=2.55):
    """Include / extend: dashed shaft, open arrowhead at p2."""
    x1, y1 = p1
    x2, y2 = p2
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1e-6)
    ux, uy = dx / length, dy / length
    # Sit the tip just off the target outline.
    clearance = 0.35
    x2 -= ux * clearance
    y2 -= uy * clearance
    end = (x2 - ux * head, y2 - uy * head)
    ax.plot(
        [x1, end[0]],
        [y1, end[1]],
        color=color,
        lw=lw,
        linestyle=(0, (5.5, 3.0)),
        zorder=2,
        solid_capstyle="butt",
    )
    # Fill at most one dash gap so the shaft meets the back of the chevron.
    ax.figure.canvas.draw()
    origin = ax.transData.transform((0.0, 0.0))
    unit = ax.transData.transform((ux, uy))
    pts_per_data = max(((unit[0] - origin[0]) ** 2 + (unit[1] - origin[1]) ** 2) ** 0.5, 1e-6)
    stub = 4.2 / pts_per_data
    ax.plot(
        [end[0] - ux * stub, end[0]],
        [end[1] - uy * stub, end[1]],
        color=color,
        lw=lw,
        solid_capstyle="butt",
        zorder=4,
    )
    open_head(ax, (x1, y1), (x2, y2), size=head, color=color, lw=lw)


def flow_label(ax, x, y, s, size=8):
    text(ax, x, y, s, size=size, color=MUTED, style="italic")


def poly_arrow(ax, pts, label=None, lxy=None, lsize=8):
    xs = [p[0] for p in pts]
    ys = [p[1] for p in pts]
    ax.plot(xs, ys, color=NAVY, lw=1.25, zorder=2, solid_capstyle="round")
    ax.add_patch(
        FancyArrowPatch(
            pts[-2],
            pts[-1],
            arrowstyle="-|>",
            mutation_scale=11,
            linewidth=1.25,
            color=NAVY,
            zorder=2,
            shrinkA=0,
            shrinkB=0,
        )
    )
    if label and lxy:
        flow_label(ax, lxy[0], lxy[1], label, size=lsize)


def stick(ax, x, y, label, scale=1.0):
    r = 2.15 * scale
    ax.add_patch(Circle((x, y + 11.2 * scale), r, fill=False, edgecolor=NAVY, lw=1.7, zorder=4))
    ax.plot([x, x], [y + 9.0 * scale, y + 3.6 * scale], color=NAVY, lw=1.7, zorder=4)
    ax.plot([x - 3.3 * scale, x + 3.3 * scale], [y + 7.4 * scale, y + 7.4 * scale], color=NAVY, lw=1.7, zorder=4)
    ax.plot([x, x - 2.5 * scale], [y + 3.6 * scale, y + 0.2 * scale], color=NAVY, lw=1.7, zorder=4)
    ax.plot([x, x + 2.5 * scale], [y + 3.6 * scale, y + 0.2 * scale], color=NAVY, lw=1.7, zorder=4)
    text(ax, x, y - 2.2, label, size=9.5, weight="bold")


def oval(ax, cx, cy, w, h, name, fill=WHITE):
    e = Ellipse((cx, cy), w, h, facecolor=fill, edgecolor=NAVY, linewidth=1.55, zorder=3)
    ax.add_patch(e)
    text(ax, cx, cy, name, size=9)


def system_box(ax, x, y, w, h, title):
    rect(ax, x, y, w, h, facecolor=WHITE, lw=1.7, z=1)
    text(ax, x + w / 2, y + h + 2.4, title, size=12, weight="bold", color=TITLE_C)


# ---------------------------------------------------------------------------
# Diagrams
# ---------------------------------------------------------------------------

def draw_uc1():
    fig, ax = new_fig(11.4, 6.4)
    system_box(ax, 22, 8, 74, 82, "Dungeon Crawler Carl — Loot Drops")
    stick(ax, 10, 38, "Map\nManager", scale=1.05)
    oval(ax, 52, 62, 28, 16, "Roll Drop\nfor Room")
    oval(ax, 38, 24, 26, 14, "Select Weighted\nTemplate")
    oval(ax, 70, 24, 26, 14, "Clone Item\nPrototype")
    oval(ax, 84, 74, 24, 14, "Guarantee\nRequired Drop")
    assoc(ax, (14.5, 49.5), (38.2, 62))
    dashed(ax, (50, 54), (40, 31.2))
    dashed(ax, (54, 54), (68, 31.2))
    dashed(ax, (73.5, 71.5), (65.2, 67.2))
    flow_label(ax, 42, 42, "<< include >>")
    flow_label(ax, 66, 42, "<< include >>")
    flow_label(ax, 76, 83.5, "<< extend >>")
    return save(fig, "uc1_roll_drop.png")


def draw_uc2():
    fig, ax = new_fig(11.0, 6.2)
    system_box(ax, 24, 10, 70, 78, "Dungeon Crawler Carl — Item Application")
    stick(ax, 11, 40, "Player", scale=1.05)
    oval(ax, 46, 68, 24, 14, "Collect Item")
    oval(ax, 46, 36, 24, 14, "Apply Effect")
    oval(ax, 80, 36, 22, 14, "Revert\nConsumable")
    oval(ax, 80, 68, 22, 14, "Restore\nLoadout")
    assoc(ax, (15.2, 51.5), (34.2, 68))
    dashed(ax, (46, 61), (46, 43.2))
    dashed(ax, (69, 36), (58.2, 36))
    dashed(ax, (69, 68), (58.2, 68))
    flow_label(ax, 52.8, 52.2, "<< include >>")
    flow_label(ax, 64, 41.2, "<< extend >>")
    flow_label(ax, 64, 73.2, "<< extend >>")
    return save(fig, "uc2_apply_item.png")


def draw_dfd0():
    """One Gane–Sarson process per Feature Specifications entry, plus shared stores."""
    fig, ax = new_fig(14.8, 8.2, equal=False, xmax=122)
    text(ax, 59, 97.4, "Diagram 0 — Dungeon Crawler Carl", size=12.5, weight="bold", color=TITLE_C)

    # Same visual grammar as the sample: numbered processes, Init-style forks,
    # CheckForCollision-style links, OnExitScene into the HUD, store under its owner.
    gane_process(ax, 16, 52, 24, 13, "2", "Generate Dungeon Map")
    gane_process(ax, 44, 84, 24, 13, "3", "Spawn Encounters")
    gane_process(ax, 76, 84, 24, 13, "6", "Run Boss Encounter")
    gane_process(ax, 104, 84, 24, 13, "1", "Update HUD & State")
    gane_process(ax, 76, 52, 24, 13, "4", "Upgrade & Loot", fill=PEACH)
    gane_process(ax, 104, 52, 24, 13, "5", "Track Points & Unlocks")
    gane_process(ax, 44, 20, 24, 13, "7", "Control Player & Combat")

    data_store(ax, 16, 20, 24, 9, "D4", "Dungeon Layout")
    data_store(ax, 44, 4, 24, 8, "D3", "Player Stats")
    data_store(ax, 76, 20, 24, 9, "D1", "Item Catalog")
    data_store(ax, 76, 4, 24, 8, "D2", "Player Loadout")
    data_store(ax, 104, 20, 24, 9, "D5", "Point Balance")

    # 2 → 3 / 7  are the Init fork. 3 → 4 and 6 → 1 match CheckForCollision / OnExitScene.
    poly_arrow(ax, [(104, 90.5), (104, 93.5), (16, 93.5), (16, 58.5)], "start run", (8.4, 76), 7.5)
    poly_arrow(ax, [(28, 57), (32, 77.5)], "room + tier", (22.6, 70.5), 7.3)
    poly_arrow(ax, [(28, 47), (32, 26.5)], "room bounds", (22.8, 35.8), 7.3)
    poly_arrow(ax, [(28, 52), (64, 52)], "room type", (46, 55.2), 7.5)
    poly_arrow(ax, [(28, 58.5), (64, 77.5)], "BossRoom", (40.5, 72.6), 7.3)
    poly_arrow(ax, [(56, 77.5), (64, 58.5)], "room cleared", (66.8, 70.4), 7.3)
    poly_arrow(ax, [(56, 84), (64, 84)], "summon minions", (60, 87.2), 7.2)
    poly_arrow(ax, [(88, 84), (92, 84)], "OnBossDefeated", (90, 87.2), 7.2)
    poly_arrow(ax, [(64, 52), (56, 26.5)], "stat deltas", (52.4, 41.5), 7.3)
    poly_arrow(ax, [(56, 13.5), (56, 11), (119, 11), (119, 84), (116, 84)], "GetStats", (119.6, 48), 7.3)
    poly_arrow(ax, [(104, 77.5), (104, 58.5)], "CreditRunEnd", (110.8, 69.2), 7.3)
    poly_arrow(ax, [(92, 52), (88, 52)], "unlocks", (90, 55.2), 7.5)

    # One link per store, like Current Score under Update Score.
    poly_arrow(ax, [(16, 45.5), (16, 24.5)], "run layout", (9.6, 35), 7.2)
    poly_arrow(ax, [(44, 13.5), (44, 8)], "combat stats", (37.2, 10.8), 7.0)
    poly_arrow(ax, [(76, 45.5), (76, 24.5)], "templates, weights", (66.4, 35), 7.0)
    poly_arrow(ax, [(88, 45.5), (90, 45.5), (90, 4), (88, 4)], "updated loadout", (97.2, 8.4), 7.0)
    poly_arrow(ax, [(104, 45.5), (104, 24.5)], "balance", (110.8, 35), 7.2)

    return save(fig, "dfd0_context.png")


def draw_dfd4():
    """Feature 4 zoom comes from the standalone matplotlib Diagram 4."""
    from build_dfd4 import draw_dfd4 as _draw
    return _draw()


def draw_tree():
    fig, ax = new_fig(12.6, 7.2, equal=False)
    text(
        ax,
        50,
        97.5,
        "Process Description — 4.2 Apply Effect (decision tree)",
        size=12.5,
        weight="bold",
        color=TITLE_C,
    )

    def dbox(cx, cy, w, h, label):
        rounded(ax, cx - w / 2, cy - h / 2, w, h, facecolor=PEACH, radius=0.16)
        text(ax, cx, cy, label, size=8.5)

    def tbox(cx, cy, w, h, label):
        rounded(ax, cx - w / 2, cy - h / 2, w, h, facecolor=GRAY, radius=0.16)
        text(ax, cx, cy, label, size=8.2)

    dbox(16, 50, 22, 13, "Item.Kind?")
    dbox(46, 82, 22, 12, "Already owned\nRelicId?")
    dbox(46, 50, 22, 12, "Same consumable\nactive?")
    tbox(78, 90, 24, 11, "Relic: MoveSpeed +0.15\nCritChance +10")
    tbox(78, 74, 24, 11, "Reject duplicate\nstats unchanged")
    tbox(78, 56, 24, 11, "Refresh duration\nto 10.0 s; no stack")
    tbox(78, 40, 24, 11, "Attack +8 for 10.0 s")
    tbox(46, 30, 22, 11, "Weapon: Attack +5")
    tbox(46, 16, 22, 11, "Armor: Defense +3\nMaxHP +10, HP +10")
    tbox(46, 4, 22, 10, "None / NullItem: no-op")

    poly_arrow(ax, [(21, 56.5), (21, 82), (35, 82)], "Relic", (22.2, 70), 8)
    poly_arrow(ax, [(27, 50), (35, 50)], "Consumable", (31, 53.6), 7.5)
    poly_arrow(ax, [(21, 43.5), (21, 30), (35, 30)], "Weapon", (22.4, 37.2), 8)
    poly_arrow(ax, [(16, 43.5), (16, 16), (35, 16)], "Armor", (17.4, 24), 8)
    poly_arrow(ax, [(16, 43.5), (16, 4), (35, 4)], "None", (10.6, 10), 8)

    poly_arrow(ax, [(57, 86), (66, 90)], "no", (60, 90.6), 8)
    poly_arrow(ax, [(57, 82), (66, 74)], "yes", (60, 79.4), 8)
    poly_arrow(ax, [(57, 54), (66, 56)], "yes", (60, 58.4), 8)
    poly_arrow(ax, [(57, 46), (66, 40)], "no", (60, 42), 8)

    return save(fig, "proc_4_2_tree.png")


def pert_node(ax, cx, cy, num, es, dur, ef, ls, slack, lf, critical=False):
    w, h = 14.2, 16.5
    x, y = cx - w / 2, cy - h / 2
    fill = PEACH if critical else WHITE
    lw = 2.0 if critical else 1.35
    rect(ax, x, y, w, h, facecolor=fill, lw=lw)
    col_w, row_h = w / 3, h / 3
    for i in range(1, 3):
        ax.plot([x + i * col_w, x + i * col_w], [y, y + row_h], color=NAVY, lw=0.9, zorder=4)
        ax.plot([x + i * col_w, x + i * col_w], [y + 2 * row_h, y + h], color=NAVY, lw=0.9, zorder=4)
    ax.plot([x, x + w], [y + row_h, y + row_h], color=NAVY, lw=0.9, zorder=4)
    ax.plot([x, x + w], [y + 2 * row_h, y + 2 * row_h], color=NAVY, lw=0.9, zorder=4)
    vals_top = [es, dur, ef]
    vals_bot = [ls, slack, lf]
    for i, v in enumerate(vals_top):
        text(ax, x + col_w * (i + 0.5), y + h - row_h / 2, str(v), size=8)
    text(ax, cx, cy, str(num), size=13, weight="bold")
    for i, v in enumerate(vals_bot):
        text(ax, x + col_w * (i + 0.5), y + row_h / 2, str(v), size=8)


def draw_pert():
    fig, ax = new_fig(13.6, 7.8, equal=False)
    text(
        ax,
        50,
        97.2,
        "PERT — Feature 4 work items only (Upgrade & Loot; hours)",
        size=12,
        weight="bold",
        color=TITLE_C,
    )

    names = {
        1: "Item hierarchy",
        2: "Loot tables",
        3: "Clone()",
        4: "Weapon / Armor",
        5: "Consumable / Relic",
        6: "RollDrop",
        7: "ApplyItem",
        8: "Memento",
        9: "F4 unit tests",
        10: "F4 docs",
    }

    # (num, es, dur, ef, ls, slack, lf, critical, x, y)
    nodes = [
        (1, 0, 5, 5, 0, 0, 5, True, 10, 60),
        (2, 0, 5, 5, 10, 10, 15, False, 10, 24),
        (3, 5, 4, 9, 5, 0, 9, True, 28, 60),
        (4, 9, 5, 14, 9, 0, 14, True, 46, 76),
        (5, 9, 5, 14, 9, 0, 14, True, 46, 44),
        (6, 9, 6, 15, 15, 6, 21, False, 28, 24),
        (7, 14, 4, 18, 14, 0, 18, True, 64, 60),
        (8, 18, 3, 21, 18, 0, 21, True, 80, 60),
        (9, 21, 10, 31, 21, 0, 31, True, 94, 76),
        (10, 21, 5, 26, 26, 5, 31, False, 94, 34),
    ]
    for n in nodes:
        pert_node(ax, n[8], n[9], n[0], n[1], n[2], n[3], n[4], n[5], n[6], n[7])
        text(ax, n[8], n[9] - 10.6, names[n[0]], size=7.2, color=MUTED)

    poly_arrow(ax, [(17.1, 60), (20.9, 60)])
    poly_arrow(ax, [(35.1, 64), (39.0, 76)])
    poly_arrow(ax, [(35.1, 56), (39.0, 44)])
    poly_arrow(ax, [(53.1, 76), (57.0, 64)])
    poly_arrow(ax, [(53.1, 44), (57.0, 56)])
    poly_arrow(ax, [(71.1, 60), (72.9, 60)])
    poly_arrow(ax, [(87.1, 64), (86.9, 76)])
    poly_arrow(ax, [(87.1, 56), (86.9, 42)])
    poly_arrow(ax, [(17.1, 24), (20.9, 24)])
    poly_arrow(ax, [(35.1, 24), (80, 24), (80, 34), (86.9, 34)])
    poly_arrow(ax, [(28, 32.3), (37, 32.3), (37, 92), (94, 92), (94, 84.3)])

    text(
        ax,
        50,
        3.8,
        "Node IDs are Feature 4 tasks (not Diagram 0 feature numbers).  Top: ES | duration | EF     Bottom: LS | slack | LF",
        size=8.2,
        color=MUTED,
    )
    return save(fig, "pert.png")


def draw_gantt():
    fig, ax = plt.subplots(figsize=(12.6, 6.8))
    fig.patch.set_facecolor("white")
    tasks = [
        ("1. Item hierarchy & effect matrix", 0, 5, 0),
        ("2. Loot table design (per RoomType)", 0, 5, 10),
        ("3. Abstract Item + Prototype Clone()", 5, 4, 0),
        ("4. WeaponUpgrade + ArmorUpgrade", 9, 5, 0),
        ("5. ConsumableItem + RelicItem", 9, 5, 0),
        ("6. LootTable.RollDrop(RoomType)", 9, 6, 6),
        ("7. ApplyItem / revert on PlayerStats", 14, 4, 0),
        ("8. ItemMemento SaveState/RestoreState", 18, 3, 0),
        ("9. Unit tests (deltas, drops, restore)", 21, 10, 0),
        ("10. Feature 4 documentation", 21, 5, 5),
    ]
    ax.set_xlim(0, 32)
    ax.set_ylim(-0.7, 10.6)
    ax.invert_yaxis()
    ax.set_xlabel("Feature 4 hour", fontsize=10, color=TEXT)
    ax.set_yticks(range(len(tasks)))
    ax.set_yticklabels([t[0] for t in tasks], fontsize=9)
    ax.set_xticks(range(0, 33, 2))
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color(NAVY)
    ax.spines["bottom"].set_color(NAVY)
    ax.tick_params(colors=TEXT)
    ax.xaxis.grid(True, linestyle=":", color="#D0D5DA")
    ax.set_axisbelow(True)
    ax.set_title(
        "Gantt — Feature 4 Upgrade & Loot only (dark = my work, light = slack)",
        fontsize=13,
        fontweight="bold",
        color=TITLE_C,
        pad=12,
    )
    bar_h = 0.55
    for i, (_, start, dur, slack) in enumerate(tasks):
        ax.barh(i, dur, left=start, height=bar_h, color=GOLD, edgecolor="#8A6A14", linewidth=0.6, zorder=3)
        if slack:
            ax.barh(
                i,
                slack,
                left=start + dur,
                height=bar_h,
                color=GOLD_LIGHT,
                edgecolor="#D9C79A",
                linewidth=0.5,
                zorder=2,
            )
    fig.tight_layout()
    path = IMG / "gantt.png"
    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white", pad_inches=0.18)
    plt.close(fig)
    return path


# ---------------------------------------------------------------------------
# Word document
# ---------------------------------------------------------------------------

def set_run_font(run, name, size, bold=False, color=None, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill: str):
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_text(cell, text_value, *, bold=False, color=None, size=10, fill=None, align="left"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text_value)
    set_run_font(run, "Calibri", size, bold=bold, color=color)
    set_cell_margins(cell)
    if fill:
        shade_cell(cell, fill)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), "9676")
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=RGBColor(255, 255, 255), fill="1E2761", size=10)
    for r, row in enumerate(rows, start=1):
        fill = "F7F4EE" if r % 2 == 0 else None
        for c, val in enumerate(row):
            set_cell_text(table.rows[r].cells[c], val, fill=fill, size=10)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table


def add_body(doc, text_value, *, first_line=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text_value)
    set_run_font(run, "Calibri", 11)
    return p


def add_label(doc, label, value=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(2)
    r1 = p.add_run(label)
    set_run_font(r1, "Calibri", 11, bold=True)
    if value:
        r2 = p.add_run(value)
        set_run_font(r2, "Calibri", 11)
    return p


def add_heading(doc, text_value):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text_value)
    set_run_font(run, "Cambria", 15, bold=True, color=NAVY_RGB)
    return p


def add_sub(doc, text_value):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text_value)
    set_run_font(run, "Cambria", 12.5, bold=True, color=NAVY_RGB)
    return p


def add_picture(doc, path: Path, width=6.45):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    return p


def build_doc(images: dict):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.875)
    section.right_margin = Inches(0.875)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)

    t = doc.add_paragraph()
    t.paragraph_format.space_after = Pt(10)
    r = t.add_run("Champion Document — Upgrade & Loot System")
    set_run_font(r, "Cambria", 26, bold=True, color=NAVY_RGB)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(10)
    mr = meta.add_run(
        "Name: ____________________        Feature: Upgrade & Loot System (Dungeon Crawler Carl)        Mark: ____ / 30"
    )
    set_run_font(mr, "Calibri", 11)

    note = doc.add_paragraph()
    nr = note.add_run(
        "This Champion covers Feature 4 (TL3). Exceptions reference the exact basic-sequence steps they branch from, "
        "DFDs use Gane–Sarson process rectangles and data stores, and every interface entry is typed with numeric pass/fail bounds."
    )
    set_run_font(nr, "Calibri", 11, italic=True)

    # 1. Introduction
    add_heading(doc, "1.  Brief Introduction")
    add_body(
        doc,
        "My feature for Dungeon Crawler Carl is Feature 4 — the upgrade and loot system (TL3). When a room becomes "
        "eligible for loot, my code looks up the LootTable for that RoomType, rolls a drop, and clones an Item prototype "
        "rather than constructing a new definition from scratch (GoF Prototype). The internal classes are abstract Item "
        "(non-MonoBehaviour), WeaponUpgradeItem, ArmorUpgradeItem, ConsumableItem, RelicItem, LootTable, and ItemMemento. "
        "Each item subclass overrides virtual void ApplyEffect(PlayerStats stats). If that virtual is removed, every "
        "pickup applies the no-op default and the HUD stops changing — the spec’s visible failure mode. Item is the "
        "Information Expert for its own effect; loot logic stays out of combat and map code (GRASP High Cohesion).",
    )
    add_body(
        doc,
        "The spec public interface is exactly four methods: ApplyItem(Item, PlayerStats), RollDrop(RoomType), "
        "SaveState() → ItemSnapshot, and RestoreState(ItemSnapshot). ApplyItem writes documented deltas into the live "
        "PlayerStats instance owned by Feature 7. ConsumableItem is temporary: Attack +8 for exactly 10.0 seconds, then "
        "the effect reverts. Weapon, armor, and relic effects last for the run unless RestoreState replays an "
        "ItemMemento (GoF Memento) taken at a checkpoint or before death. SaveState captures equipped item ids, remaining "
        "consumable timers, and the stat baseline so a round-trip restore is equal on every numeric field.",
    )
    add_body(
        doc,
        "My feature deliberately does not own combat, movement, ApplyDamage, GetPosition, or the PlayerStats type — those "
        "belong to Feature 7; I only mutate a PlayerStats instance I am given. It does not generate runs, pick room "
        "templates, or implement TryTransition — Feature 2’s MapManager / RoomFactory own CombatRoom, TreasureRoom, "
        "CorridorRoom, and BossRoom; I only accept that RoomType as RollDrop input. It does not spawn encounters or "
        "raise OnEnemyDefeated (Feature 3). It does not run boss phases or OnBossDefeated (Feature 6); BossRoom is only "
        "a drop-weighting key. It does not render HUD elements or broadcast GameState — Feature 1 observes "
        "GameStateManager.OnStateChanged and reads stats via Feature 7 GetStats() after my mutation. It does not "
        "implement EconomyManager, TryPurchase, or point balance (Feature 5). Feature 5’s spec says purchases may later "
        "gate Feature 4 templates; that unlock set is not on my public interface and is out of scope for this Champion.",
    )

    # 2. Use Case
    add_heading(doc, "2.  Use Case Diagram with Scenario")
    add_sub(doc, "Use Case Diagrams")
    add_picture(doc, images["uc1"])
    add_picture(doc, images["uc2"])

    add_sub(doc, "Scenarios")
    add_sub(doc, "Scenario 1 (first use case diagram)")
    add_label(doc, "Name: ", "Roll Drop for Room")
    add_label(
        doc,
        "Summary: ",
        "When a room becomes eligible for loot, Feature 2’s MapManager (or Feature 3/6 passing the current room’s type after a clear) requests a drop; a weighted template is selected from that RoomType’s LootTable and cloned via Prototype.",
    )
    add_label(doc, "Actors: ", "MapManager (Feature 2); Feature 3 EncounterSpawner / Feature 6 BossController may call the same method with the current RoomType")
    add_label(
        doc,
        "Preconditions: ",
        "The Item catalog and per-RoomType LootTable entries are loaded; RoomType is one of Feature 2’s four kinds: CombatRoom, TreasureRoom, CorridorRoom, BossRoom.",
    )
    add_label(doc, "Basic sequence:")
    add_body(doc, "Step 1: The caller passes RoomType (and the run seed used by the loot RNG).")
    add_body(doc, "Step 2: LootTable for that RoomType is retrieved from the catalog.")
    add_body(doc, "Step 3: A drop-chance roll in [0.0, 1.0) is compared to the room’s drop rate.")
    add_body(doc, "Step 4: A weighted template is selected from the table (weights are positive integers that sum to 100).")
    add_body(doc, "Step 5: The template is cloned with Item.Clone() into a new Item instance (Prototype).")
    add_body(doc, "Step 6: RollDrop returns that Item (Null Object if the chance roll failed).")
    add_label(doc, "Exceptions:")
    add_body(
        doc,
        "Step 2a: RoomType has no table. RollDrop throws InvalidRoomTypeException and does not continue. Unit tests "
        "assert the exception type; production map code never passes an undefined enum value.",
    )
    add_body(
        doc,
        "Step 3a: The drop-chance roll is ≥ the room’s drop rate (CombatRoom 0.35, CorridorRoom 0.10). Sequence skips to Step 6 "
        "with NullItem (ItemId = \"NULL\", Kind = None) — a Null Object implementing Item, not a fifth item subclass. "
        "TreasureRoom and BossRoom never take this branch (drop rate 1.00).",
    )
    add_body(
        doc,
        "Step 4a: Weighted selection yields an unknown template id — the <<extend>> Guarantee Required Drop case "
        "for TreasureRoom and BossRoom only. Selection is retried up to 5 times; if still invalid, the test harness fails. "
        "CombatRoom/CorridorRoom return NullItem after one failure and log a warning. The sequence resumes at Step 6.",
    )
    add_label(doc, "Post conditions: ", "The caller holds either a cloned Item whose Kind is one of the four types, or NullItem. Catalog templates are unchanged.")
    add_label(doc, "Priority: ", "1 (must have)")
    add_label(doc, "ID: ", "LD1")

    add_sub(doc, "Scenario 2 (second use case diagram)")
    add_label(doc, "Name: ", "Collect Item")
    add_label(
        doc,
        "Summary: ",
        "The player collects a dropped item; ApplyItem mutates PlayerStats by the item’s documented deltas. Consumables later revert; death/checkpoint restores an ItemMemento.",
    )
    add_label(doc, "Actors: ", "Player")
    add_label(
        doc,
        "Preconditions: ",
        "A non-null Item exists to collect (or RestoreState is invoked with a previously saved ItemSnapshot); PlayerStats is the live Feature 7 instance.",
    )
    add_label(doc, "Basic sequence:")
    add_body(doc, "Step 1: The player collects the item (pickup input).")
    add_body(doc, "Step 2: ApplyItem(item, stats) is called with non-null arguments.")
    add_body(doc, "Step 3: item.ApplyEffect(stats) applies the typed delta for that subclass.")
    add_body(doc, "Step 4: Permanent items are appended to the loadout; a ConsumableItem starts a 10.0 s timer.")
    add_body(doc, "Step 5: ApplyItem returns; Feature 7 GetStats() now reflects the new fields. Feature 1 HUD updates via GameStateManager.OnStateChanged — it does not subscribe to Feature 4.")
    add_body(doc, "Step 6: SaveState() writes an ItemSnapshot into ItemMemento (equipped ids, timers, stat baseline).")
    add_label(doc, "Exceptions:")
    add_body(doc, "Step 2a: item is NullItem. ApplyItem is a no-op: no loadout change, every PlayerStats field unchanged.")
    add_body(
        doc,
        "Step 3a: RelicItem whose RelicId is already in the loadout. ApplyEffect is skipped; "
        "MoveSpeed and CritChancePercent stay at their pre-call values.",
    )
    add_body(
        doc,
        "Step 4a: A ConsumableItem with the same ConsumableId is already active — the <<extend>> Revert Consumable / "
        "refresh case. Attack is not increased a second time; remaining duration is reset to 10.0 s.",
    )
    add_body(
        doc,
        "Step 6a: Player HP reaches 0 or a checkpoint fires — the <<extend>> Restore Loadout case. "
        "RestoreState(snap) replaces the live loadout with the snapshot; remaining consumable timers are restored; "
        "each PlayerStats field equals the snapshot baseline after re-applying saved items.",
    )
    add_label(doc, "Post conditions: ", "PlayerStats matches the documented deltas (or the restored snapshot). HUD-visible fields have changed only when an effect actually applied.")
    add_label(doc, "Priority: ", "1 (must have)")
    add_label(doc, "ID: ", "CI1")

    # 3. Interface Contract
    add_heading(doc, "3.  Interface Contract")
    add_body(
        doc,
        "The spec public interface is the four methods below — not events. Feature 1 already uses Observer on "
        "GameStateManager; the HUD must not subscribe here. Every input is an enum, a named type, or a numeric field. "
        "Every output is asserted in headless unit tests against the exact bounds in the tables below. No scene is required.",
    )
    add_body(
        doc,
        "void ApplyItem(Item item, PlayerStats stats)     Item RollDrop(RoomType roomType)     "
        "ItemSnapshot SaveState()     void RestoreState(ItemSnapshot snap)",
    )

    add_sub(doc, "Types")
    add_body(
        doc,
        "RoomType matches Feature 2’s four room kinds: enum { CombatRoom, TreasureRoom, CorridorRoom, BossRoom }. "
        "ItemKind: enum { None, WeaponUpgrade, ArmorUpgrade, Consumable, Relic }. "
        "PlayerStats is Feature 7’s type; Feature 4 requires these fields and run defaults: int MaxHealth = 100, "
        "int CurrentHealth = 100, int Attack = 10, int Defense = 0, float MoveSpeed = 1.00f, int CritChancePercent = 0. "
        "ItemSnapshot: EquippedItemIds : ItemId[], ConsumableTimers : (ItemId, RemainingSeconds : float)[], "
        "StatBaseline : PlayerStats, SchemaVersion : int = 1.",
    )

    add_sub(doc, "Inputs")
    add_table(
        doc,
        ["Input", "Type", "Accepted values / bounds", "On violation"],
        [
            ["ApplyItem.item", "Item", "Non-null. NullItem allowed (no-op).", "ArgumentNullException if null"],
            ["ApplyItem.stats", "PlayerStats", "Non-null live stats instance.", "ArgumentNullException if null"],
            ["RollDrop.roomType", "RoomType", "CombatRoom | TreasureRoom | CorridorRoom | BossRoom", "InvalidRoomTypeException"],
            ["RestoreState.snap", "ItemSnapshot", "Non-null; SchemaVersion == 1", "ArgumentNullException / SnapshotVersionException"],
            ["Loot RNG seed", "int", "Same seed + table version ⇒ identical ItemId sequence", "Test fail if two runs diverge"],
        ],
        col_widths=[1.6, 1.3, 2.4, 1.5],
    )

    add_sub(doc, "Outputs")
    add_table(
        doc,
        ["Output", "Type", "Pass bound", "Fail if"],
        [
            ["RollDrop → CombatRoom", "Item", "NullItem with p = 0.65 ± 0.04 over 2,000 rolls; else one of WEAPON_COMMON, ARMOR_COMMON, CONSUMABLE_RAGE", "Any other ItemId, or NullItem rate outside 0.61–0.69"],
            ["RollDrop → TreasureRoom", "Item", "Never NullItem. ItemId ∈ {WEAPON_RARE, ARMOR_RARE, RELIC_SWIFT, RELIC_LUCKY}", "NullItem or unknown id"],
            ["RollDrop → CorridorRoom", "Item", "NullItem with p = 0.90 ± 0.03 over 2,000 rolls; else CONSUMABLE_RAGE or ARMOR_COMMON", "NullItem rate outside 0.87–0.93"],
            ["RollDrop → BossRoom", "Item", "Never NullItem. ItemId ∈ {RELIC_SWIFT, RELIC_LUCKY}", "NullItem or non-relic"],
            ["WeaponUpgrade Apply", "void / stats", "Attack += 5; all other fields unchanged", "Any other field differs"],
            ["ArmorUpgrade Apply", "void / stats", "Defense += 3; MaxHealth += 10; CurrentHealth += 10, then clamp to MaxHealth", "Attack/MoveSpeed/Crit change, or HP not clamped"],
            ["Consumable Apply", "void / stats", "Attack += 8; DurationRemaining = 10.0 ± 0.0 s at t=0", "Stack beyond +8, or duration ≠ 10.0"],
            ["Consumable Revert", "void / stats", "At t = 10.0 s, Attack returns to pre-apply value (± 0)", "Revert early (< 9.99 s) or late with Attack still buffed at 10.05 s"],
            ["Relic Apply", "void / stats", "MoveSpeed += 0.15f (± 0.0001); CritChancePercent += 10; duplicate RelicId is a no-op", "Second copy changes stats"],
            ["SaveState / RestoreState", "ItemSnapshot", "RestoreState(SaveState()) ⇒ every PlayerStats field equal; timers within 0.01 s", "Any field mismatch"],
            ["Item.Clone()", "Item", "Clone.Kind == template.Kind; clone is not reference-equal to template; mutating clone does not mutate template", "Shared mutation or Kind mismatch"],
        ],
        col_widths=[1.7, 1.2, 2.5, 1.4],
    )

    add_sub(doc, "Test-only notifications (not on the spec public interface)")
    add_body(
        doc,
        "Feature 4’s published API is the four methods. These notifications exist so headless tests can observe side effects "
        "without a scene. Feature 1 must keep using GameStateManager.OnStateChanged; it must not subscribe here.",
    )
    add_table(
        doc,
        ["Notification", "Payload type", "Raised when", "Testable assertion"],
        [
            ["OnDropRolled", "DropRolledEvent { RoomType roomType; string itemId; int rngSeed }", "End of every RollDrop call, including NullItem", "Exactly 1 notification per call; itemId matches return value"],
            ["OnItemApplied", "ItemAppliedEvent { string itemId; ItemKind kind; StatDelta delta }", "ApplyEffect actually mutated stats", "Not raised for NullItem or duplicate relic"],
            ["OnItemReverted", "ItemRevertedEvent { string itemId; StatDelta delta }", "Consumable duration reaches 0.0 s", "Raised once; Attack restored; not raised for permanent items"],
            ["OnLoadoutRestored", "ItemSnapshot", "RestoreState completes", "Payload equals the snap argument; stats match"],
        ],
        col_widths=[1.5, 2.2, 1.6, 1.5],
    )

    add_sub(doc, "Loot table weights (sum to 100 per RoomType)")
    add_table(
        doc,
        ["RoomType", "Drop rate", "Template id", "Weight"],
        [
            ["CombatRoom", "0.35", "WEAPON_COMMON", "40"],
            ["CombatRoom", "0.35", "ARMOR_COMMON", "40"],
            ["CombatRoom", "0.35", "CONSUMABLE_RAGE", "20"],
            ["TreasureRoom", "1.00", "WEAPON_RARE", "25"],
            ["TreasureRoom", "1.00", "ARMOR_RARE", "25"],
            ["TreasureRoom", "1.00", "RELIC_SWIFT", "25"],
            ["TreasureRoom", "1.00", "RELIC_LUCKY", "25"],
            ["CorridorRoom", "0.10", "CONSUMABLE_RAGE", "70"],
            ["CorridorRoom", "0.10", "ARMOR_COMMON", "30"],
            ["BossRoom", "1.00", "RELIC_SWIFT", "50"],
            ["BossRoom", "1.00", "RELIC_LUCKY", "50"],
        ],
        col_widths=[1.5, 1.3, 2.2, 1.2],
    )

    add_body(
        doc,
        "StatDelta is { int attack; int defense; int maxHealth; int currentHealth; float moveSpeed; int critChancePercent }. "
        "COMMON templates use the deltas in the table; RARE weapon/armor templates are the same subclasses with AttackDelta = 8, "
        "DefenseDelta = 5, MaxHealthDelta = 20. Unit tests apply each of the four item types (common and rare where they exist) "
        "to a mock PlayerStats and assert the exact field-by-field delta. Distribution tests use a fixed seed and 2,000 rolls "
        "per RoomType. Memento tests save, mutate stats with a second item, restore, and require equality with the snapshot.",
    )

    # 4. DFD
    add_heading(doc, "4.  Data Flow Diagrams")
    add_body(
        doc,
        "Diagram 0 shows the whole game: one numbered process for each feature in the Feature Specifications "
        "(1 HUD & game state, 2 dungeon map, 3 encounters, 4 upgrade & loot, 5 point economy, 6 boss, 7 player combat). "
        "Process 4 is highlighted as this Champion’s feature. I then zoom into Process 4 and give a decision-tree "
        "description for primitive process 4.2. Processes are Gane–Sarson rounded rectangles with a numbered band; "
        "data stores are open rectangles with a left-hand identifier. Every process has at least one incoming and one "
        "outgoing flow (no black holes or miracles). D1 and D2 belong to Process 4 (catalog and loadout). D3 Player Stats, "
        "D4 Dungeon Layout, and D5 Point Balance are the stores the other features need.",
    )
    add_picture(doc, images["dfd0"], width=6.8)
    add_picture(doc, images["dfd4"], width=6.8)
    add_sub(doc, "Process Description — 4.2 Apply Effect")
    add_picture(doc, images["tree"], width=6.5)
    add_body(
        doc,
        "4.2 is a primitive: it does not decompose further. It reads Item.Kind and the current loadout from D2, writes "
        "mutated PlayerStats out to Process 7, snapshots the loadout to D2, and emits an effect record to 4.3 "
        "(consumables) and 4.4 (every successful apply). "
        "A duplicate relic is a valid output of 4.2 (no-op), not a missing output. Completion criteria from the spec: "
        "all four item types apply correct, reversible effects, and every RoomType’s LootTable produces a valid Item "
        "when a drop occurs.",
    )

    # 5. Timeline
    add_heading(doc, "5.  Timeline (Feature 4 only)")
    add_body(
        doc,
        "This schedule is my Feature 4 work only — not the seven-feature project calendar. Hours match the Feature 4 "
        "spec budget: Design 10 hrs, Coding 27 hrs, Testing 10 hrs, Documentation 5 hrs (52 hrs of my work). "
        "Because several of my coding tasks run in parallel, calendar time on the critical path is 31 hours. "
        "TL4 tracks actuals against these Feature 4 rows. Features 1, 2, 3, 5, 6, and 7 keep their own champions and hour budgets.",
    )
    add_sub(doc, "Feature 4 work items")
    add_table(
        doc,
        ["Feature 4 task", "Duration (hours)", "Predecessor task(s)"],
        [
            ["1.  Item class hierarchy & effect matrix", "5", "—"],
            ["2.  Loot table design (per RoomType)", "5", "—"],
            ["3.  Abstract Item + Prototype Clone()", "4", "1"],
            ["4.  WeaponUpgradeItem + ArmorUpgradeItem", "5", "3"],
            ["5.  ConsumableItem + RelicItem", "5", "3"],
            ["6.  LootTable.RollDrop(RoomType)", "6", "2, 3"],
            ["7.  ApplyItem / revert on PlayerStats", "4", "4, 5"],
            ["8.  ItemMemento SaveState/RestoreState", "3", "7"],
            ["9.  Unit tests (deltas, drops, restore)", "10", "6, 7, 8"],
            ["10.  Feature 4 documentation", "5", "6, 8"],
        ],
        col_widths=[3.4, 1.6, 1.8],
    )
    add_sub(doc, "PERT diagram (Feature 4 tasks)")
    add_picture(doc, images["pert"], width=6.5)
    add_body(
        doc,
        "PERT nodes are the Feature 4 tasks in the table, not Diagram 0’s feature numbers. "
        "The critical path is 1 → 3 → 4 → 7 → 8 → 9 (equivalently 1 → 3 → 5 → 7 → 8 → 9), totaling 31 hours of my calendar time. "
        "Tasks 4 and 5 are both critical because either one finishing late delays ApplyItem. Task 2 carries 10 hours of "
        "slack; task 6 carries 6; documentation (task 10) carries 5. Design of the loot tables can slip without moving my finish, "
        "but the Item hierarchy cannot.",
    )
    add_sub(doc, "Gantt timeline (Feature 4 only)")
    add_picture(doc, images["gantt"], width=6.5)

    add_heading(doc, "AI Question Log — Summary")
    add_body(
        doc,
        "The full question record is docs/TL3_Kaleb/AI_Interaction_Log.md (Entries 1–8, questions 1–24). Numbers in this "
        "Champion (stat deltas, drop rates, hours, PERT values) come from the Feature 4 specification and the sample’s "
        "grading notes, not from accepting a first AI draft.",
    )
    add_body(
        doc,
        "I drafted the five required sections from the Feature 4 spec using the Enemy Spawning sample as the structural "
        "template, then checked the three things the sample warns about. First, exceptions are branches of numbered steps "
        "(Step 3a, 4a, 6a), not separate stories. Second, DFDs use Gane–Sarson process rectangles and D-numbered stores, "
        "not flowchart circles. Third, the timeline is my Feature 4 budget only — Design 10 + Coding 27 + Testing 10 + "
        "Documentation 5 = 52 hours of work / 31 hours on the critical path — not a collapsed “coding only” guess and not "
        "the seven-feature project calendar. Interface bounds (drop rates, +5 Attack, 10.0 s consumable, 0.15 MoveSpeed) "
        "are the testable contract for TL4.",
    )
    add_body(
        doc,
        "After reading all seven Feature Specifications, RoomType names match Feature 2 (CombatRoom / TreasureRoom / "
        "CorridorRoom / BossRoom), the public API is only ApplyItem, RollDrop, SaveState, and RestoreState (events are "
        "test-only), and HUD updates stay on Feature 1’s Observer. Feature 5 unlocks and Feature 6 boss phases stay out "
        "of this interface. Feature 1’s spec slip (“player stats (Feature 2)”) is not copied: GetStats() is Feature 7.",
    )
    add_body(
        doc,
        "Later passes corrected diagrams against instructor examples. Use-case lines follow the Pearson legend: "
        "Communicates is a solid line with no heads; << include >> and << extend >> are dashed with an open arrowhead "
        "(include points at the common use case; extend points from the exception to the basic use case). Diagram 0 has "
        "one process per spec feature (1–7) with process 4 highlighted; stores D1–D5 are the open-right Gane–Sarson "
        "symbol (ID band on the left). Every flow starts and ends on a box — dead-end arrows on Diagram 0 and Diagram 4 "
        "were rejected. Diagram 4 explodes process 4 (4.1 Roll Drop, 4.2 Apply Effect, 4.3 Revert, 4.4 Snapshot). PERT "
        "node numbers are Feature 4 work items, not Diagram 0 feature IDs; the Gantt axis is Feature 4 hour.",
    )

    doc.save(OUT_DOCX)
    return OUT_DOCX


def main():
    images = {
        "uc1": draw_uc1(),
        "uc2": draw_uc2(),
        "dfd0": draw_dfd0(),
        "dfd4": draw_dfd4(),
        "tree": draw_tree(),
        "pert": draw_pert(),
        "gantt": draw_gantt(),
    }
    path = build_doc(images)
    print("Wrote", path)
    for k, v in images.items():
        print(k, v, v.stat().st_size)


if __name__ == "__main__":
    main()
